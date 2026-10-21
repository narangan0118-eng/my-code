"""Select 50 distinct 10-stock portfolios per correlation tier, using training data only.
Uses pandas/openpyxl for Excel, numpy for statistics; no OOS information in selection.
"""
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

INPUT = Path(__file__).with_name('Final_Three_Portfolios_337.xlsx')
OUTPUT = Path(__file__).with_name('Robustness_150_Portfolio_Selection.xlsx')
SEED, N_TRIALS, N_PER_GROUP, K, TRAIN = 20261021, 250000, 50, 10, 120
TARGETS = {'High': .5570480544116347, 'Medium': .3036089386037136, 'Low': .101523121208463}
MAX_SHARED_WITHIN_GROUP = 6

# Read exactly the training sheet: first 120 monthly returns, excluding the date column.
if not INPUT.exists():
    raise FileNotFoundError(f"Input workbook not found: {INPUT}")
df = pd.read_excel(INPUT, sheet_name='Training_All', engine='openpyxl')
if df.shape[1] != 338:
    raise ValueError(f"Expected 337 stocks plus a date column, got {df.shape[1]} columns")
tickers = [str(c) for c in df.columns[1:]]
X = df.iloc[:TRAIN, 1:].to_numpy(dtype=float)
assert X.shape == (120, 337) and np.isfinite(X).all()
C = np.corrcoef(X,rowvar=False)
assert np.isfinite(C).all()
rng=np.random.default_rng(SEED)
# Candidate portfolios: deterministic Monte Carlo without looking at OOS returns.
choices = set()
for _ in range(N_TRIALS):
    choices.add(tuple(sorted(rng.choice(337,K,replace=False).tolist())))
scored=[]
for c in choices:
    ii=np.array(c)
    avg=float((C[np.ix_(ii,ii)].sum()-K)/(K*(K-1)))
    scored.append((avg,c))
print('candidate count',len(scored),'range',min(x[0] for x in scored),max(x[0] for x in scored))

chosen={}; used=set()
for tier,target in TARGETS.items():
    ordered=sorted(scored,key=lambda x:(abs(x[0]-target),x[1]))
    selected=[]
    for avg,c in ordered:
        if c in used: continue
        if all(len(set(c)&set(prev[1]))<=MAX_SHARED_WITHIN_GROUP for prev in selected):
            selected.append((avg,c));used.add(c)
            if len(selected)==N_PER_GROUP:break
    if len(selected)!=N_PER_GROUP:raise RuntimeError(f'Only {len(selected)} candidates for {tier}')
    chosen[tier]=selected
    print(tier, 'range',min(x[0] for x in selected),max(x[0] for x in selected),'mean',np.mean([x[0] for x in selected]))

design_rows = [
 ['Parameter','Value'], ['Source file',INPUT.name], ['Eligible stocks',337],
 ['Training months',TRAIN], ['Portfolio size',K], ['Portfolios per tier',N_PER_GROUP],
 ['Total portfolios',N_PER_GROUP*3], ['Candidate draws',N_TRIALS], ['Seed',SEED],
 ['Max shared stocks within tier',MAX_SHARED_WITHIN_GROUP],
 ['High target correlation',TARGETS['High']],
 ['Medium target correlation',TARGETS['Medium']],
 ['Low target correlation',TARGETS['Low']]
]
summary_rows = []
member_rows = []
for tier, records in chosen.items():
    for j, (avg, c) in enumerate(records, 1):
        pid = f'{tier}_{j:02d}'
        names = [tickers[i] for i in c]
        summary_rows.append({
            'Portfolio_ID': pid, 'Tier': tier,
            'Mean_Pairwise_Correlation': avg,
            'Target_Correlation': TARGETS[tier],
            'Abs_Difference': abs(avg-TARGETS[tier]),
            'Tickers': ', '.join(names)
        })
        for ticker in names:
            member_rows.append({'Portfolio_ID': pid, 'Tier': tier, 'Ticker': ticker})

with pd.ExcelWriter(OUTPUT, engine='openpyxl') as writer:
    pd.DataFrame(design_rows[1:], columns=design_rows[0]).to_excel(
        writer, sheet_name='Design', index=False)
    pd.DataFrame(summary_rows).to_excel(writer, sheet_name='Portfolio_Summary', index=False)
    pd.DataFrame(member_rows).to_excel(writer, sheet_name='Portfolio_Members', index=False)
    widths = {
        'Design': [39, 46],
        'Portfolio_Summary': [19, 14, 27, 24, 21, 95],
        'Portfolio_Members': [20, 15, 19]
    }
    for name, col_widths in widths.items():
        ws = writer.sheets[name]
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='183153')
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(wrap_text=True)
        for idx, width in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = width
    for row in writer.sheets['Portfolio_Summary'].iter_rows(min_row=2, min_col=3, max_col=5):
        for cell in row:
            cell.number_format = '0.0000'
print('SAVED', OUTPUT)
